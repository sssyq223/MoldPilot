
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from pathlib import Path

out = Path(r"E:\设计部门已完成封装成果报告.docx")
doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(0.7)
sec.bottom_margin = Inches(0.65)
sec.left_margin = Inches(0.75)
sec.right_margin = Inches(0.75)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Microsoft YaHei"
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
normal.font.size = Pt(10.5)
normal.paragraph_format.space_after = Pt(5)
normal.paragraph_format.line_spacing = 1.18
for name, size in [("Title", 22), ("Heading 1", 15), ("Heading 2", 12.5), ("Heading 3", 11)]:
    st = styles[name]
    st.font.name = "Microsoft YaHei"
    st._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = RGBColor(0,0,0)
    st.paragraph_format.space_before = Pt(10 if name != "Title" else 0)
    st.paragraph_format.space_after = Pt(5)
if "Report Subtitle" not in styles:
    st = styles.add_style("Report Subtitle", WD_STYLE_TYPE.PARAGRAPH)
    st.font.name = "Microsoft YaHei"
    st._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    st.font.size = Pt(12)
    st.font.color.rgb = RGBColor(80,80,80)
    st.paragraph_format.space_after = Pt(12)
if "Small Note" not in styles:
    st = styles.add_style("Small Note", WD_STYLE_TYPE.PARAGRAPH)
    st.font.name = "Microsoft YaHei"
    st._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    st.font.size = Pt(9)
    st.font.color.rgb = RGBColor(95,95,95)
    st.paragraph_format.space_after = Pt(4)

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcPr.append(shd)
    shd.set(qn("w:fill"), fill)

def set_cell_border(cell, color="D9D9D9", sz="4"):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = tcPr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tcPr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn("w:" + edge)
        el = borders.find(tag)
        if el is None:
            el = OxmlElement("w:" + edge)
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), sz)
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)

def set_cell_margin(cell, top=90, start=100, bottom=90, end=100):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for key, value in (("top",top),("start",start),("bottom",bottom),("end",end)):
        node = tcMar.find(qn("w:"+key))
        if node is None:
            node = OxmlElement("w:"+key)
            tcMar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")

def repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    trPr.append(OxmlElement("w:tblHeader"))

def cell_text(cell, text, bold=False, color=None, align=None, size=9.2):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    if align is not None:
        p.alignment = align
    r = p.add_run(str(text))
    r.font.name = "Microsoft YaHei"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    r.font.size = Pt(size)
    r.bold = bold
    if color:
        r.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margin(cell)
    set_cell_border(cell)

def table(headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.style = "Table Grid"
    repeat_header(t.rows[0])
    for i,h in enumerate(headers):
        cell_text(t.rows[0].cells[i], h, True, "FFFFFF", WD_ALIGN_PARAGRAPH.CENTER, 9.3)
        shade(t.rows[0].cells[i], "3F4E5E")
    for ridx,row in enumerate(rows):
        cells=t.add_row().cells
        for i,v in enumerate(row):
            cell_text(cells[i], v, align=WD_ALIGN_PARAGRAPH.CENTER if i==0 else WD_ALIGN_PARAGRAPH.LEFT)
            if ridx%2==1:
                shade(cells[i], "F5F7F9")
    if widths:
        for row in t.rows:
            for i,w in enumerate(widths):
                row.cells[i].width=Inches(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return t

def bullets(items, style="List Bullet"):
    for item in items:
        p=doc.add_paragraph(style=style)
        p.paragraph_format.space_after=Pt(2)
        p.add_run(item)

# Header
p=doc.add_paragraph(style="Title")
p.alignment=WD_ALIGN_PARAGRAPH.CENTER
p.add_run("设计部门已完成封装成果报告")
p=doc.add_paragraph(style="Report Subtitle")
p.alignment=WD_ALIGN_PARAGRAPH.CENTER
p.add_run("基于需求规格 V1.1、技术开发文档 V3.6 与当前 MoldPilot 工作树")
p=doc.add_paragraph()
p.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=p.add_run("2026年9月27日")
r.font.size=Pt(10)
r.font.color.rgb=RGBColor(100,100,100)

p=doc.add_paragraph()
p.paragraph_format.space_after=Pt(8)
r=p.add_run("本报告记录设计部门已经完成的工具封装、Skill 封装、审批材料、工作台展示和自动化验证成果。")
r.bold=True

table(
    ["已完成统计项","数量","完成内容"],
    [
        ["设计 Skill","22 个","设计路线、上传、BOM、设变、修模、标准件和订单相关 Skill"],
        ["设计相关工具","60 个","58 个 ERP 设计工具、1 个本地设计路线工具、1 个设计审批准备工具"],
        ["查询与分析工具","33 个","查询、解析、校验、核价、公差、图纸、BOM、设变和文件读取"],
        ["人工确认动作工具","25 个","正式动作统一进入人工确认卡和确认回执流程"],
        ["前端设计测试","31 项通过","覆盖上传预览、参数、公差、图纸和表单命令"],
    ],
    [1.55,1.25,4.5]
)

doc.add_heading("1 设计路线与项目上下文", level=1)
bullets([
    "已完成 query_design_route_context 设计路线查询工具。",
    "支持按项目、设计单、图纸版本、BOM 物料、计划任务和工程联络线索定位设计对象。",
    "已完成设计版本、BOM 明细、内部加工路线、采购路线和委外路线的上下文聚合。",
    "已完成关联计划任务和工程联络影响的结构化展示。",
    "已完成上一版与当前版之间的图纸版本、BOM 数量、BOM 项目、加工路线和任务关联差异分析。",
    "已完成计划复核候选的结构化输出，工作台可展示影响对象和建议动作。",
])

doc.add_heading("2 新模与改模设计清单", level=1)
bullets([
    "已完成新模钢料清单解析和新模五金清单解析。",
    "已完成改模钢料清单解析和改模五金清单解析。",
    "已完成上传会话登记、清单类型识别和设计订单类型记录。",
    "已完成图纸处理状态查询、完整上传结果查询和明细校验。",
    "已完成 ERP 钢料价格重算、热处理与时效处理字段处理。",
    "已完成钢料公差档位、长宽厚允许范围和对角公差结果封装。",
    "已完成材质、采购数量、长宽厚、料型和图纸信息查询。",
    "已完成上传会话图纸预览和标准件图纸预览入口。",
    "已完成图纸自动修正结果、修正前后值和修正原因展示。",
    "已完成新模和改模审批配置查询。",
    "已完成新模导入、改模导入和变更请购 ERP 适配器登记。",
])

doc.add_heading("3 设计订单审批材料", level=1)
bullets([
    "已完成 prepare_design_order_approval 设计订单审批准备工具。",
    "已完成 ERP 设计订单只读读取和订单明细摘要整理。",
    "已完成 ERP 来源系统、资源编号、记录版本、核对时间和快照哈希记录。",
    "已完成项目版本、设计类型、图纸版本、复核人员和审批模板信息展示。",
    "已完成当前对话附件的精确版本绑定和不可变附件记录。",
    "已完成 ERP 订单版本变化检查和过期材料阻断。",
    "已完成设计订单审批材料提交 Agent BPM 的本地业务流程。",
    "已完成审批详情中订单号、模具号、ERP 状态、版本、哈希和明细展示。",
])

doc.add_heading("4 ERP 设计工作台查询", level=1)
table(
    ["查询类别","已完成封装内容","展示结果"],
    [
        ["设计订单","订单列表、单条记录、订单状态和明细","订单号、模具号、状态、版本、明细"],
        ["图纸版本","版本列表、版本详情、版本对比","版本号、状态、时间和差异"],
        ["BOM","BOM 明细、汇总报表、采购进度、缺料","物料、数量、采购和缺料信息"],
        ["设变","设变列表、设变明细、影响分析","设变状态、影响对象和分析结果"],
        ["基础资料","材质密度、分组规则、分组关键词","ERP 原始记录和分页信息"],
        ["标准件","标准件图纸目录和文件预览","编号、文件名、目录和预览入口"],
        ["闲置料","闲置料候选、订单明细匹配和状态","材质、规格、数量和匹配状态"],
        ["修模改模","普通审批、委外审批、加工商响应","批次、订单和响应状态"],
    ],
    [1.2,3.2,2.9]
)

doc.add_heading("5 设计正式动作适配", level=1)
bullets([
    "已完成设计订单明细修改适配。",
    "已完成闲置料保存和释放适配。",
    "已完成订单删除、审批和重新提交适配。",
    "已完成材质密度、设计分组规则和分组关键词维护适配。",
    "已完成标准件图纸上传、重命名和目录删除适配。",
    "已完成 ERP 设变和设变明细维护适配。",
    "已完成修模改模数量确认、审批批次提交、订单关联和加工商响应适配。",
    "已完成 BOM 新增、修改、删除和导入适配。",
    "已完成设计文件下载和当前会话私有附件保存。",
    "已完成正式动作统一进入人工确认卡，确认后由受控确认接口调用 ERP 控制工具。",
])

doc.add_heading("6 统一工作台前端", level=1)
bullets([
    "已完成设计工具证据卡片。",
    "已完成上传结果表格和分页展示。",
    "已完成设计参数表、公差表和图纸表格。",
    "已完成 ERP 设计订单查看入口和订单明细弹窗。",
    "已完成标准件图纸目录和图纸预览入口。",
    "已完成自动修正差异的前后值和修正原因展示。",
    "已完成交期、订单类型、请购原因和备注的对话式表单编辑。",
    "已保持设计能力在统一工作台内使用，没有新增传统 ERP 式独立设计菜单。",
])

doc.add_heading("7 权限、来源和确认控制", level=1)
bullets([
    "已完成设计工具和 Skill 的设计部门归类。",
    "已完成设计能力与 design_route.read、design_route.create、design_route.execute 权限的关联。",
    "已完成 ERP 来源、业务时间和限制说明的统一返回结构。",
    "已完成当前用户上传会话和附件归属校验。",
    "已完成正式动作提案哈希校验、权限版本校验和确认凭证校验。",
    "已完成重复确认返回原有回执的幂等处理。",
    "已完成设计订单审批材料来源版本和快照哈希的审计记录。",
])

doc.add_heading("8 已通过的自动化验证", level=1)
table(
    ["验证项","结果","验证内容"],
    [
        ["设计能力目录定向测试","13 项通过","设计工具、Skill、权限元数据和人工确认模式"],
        ["前端设计预览测试","31 项通过","上传预览、参数、公差、图纸和表单命令"],
        ["Vue 类型检查","通过","mold 前端类型检查"],
        ["前端生产构建","通过","生产构建包生成成功"],
        ["确认卡门禁测试","通过","正式动作首次调用生成提案，确认后才执行"],
    ],
    [1.8,1.55,3.95]
)

doc.add_heading("9 已完成成果清单", level=1)
bullets([
    "设计部门已形成完整的 Skill 目录和 ERP 工具目录。",
    "设计查询、上传解析、清单校验、核价、公差、参数、图纸和 BOM 能力已经进入统一工作台。",
    "设计订单审批材料已经具备来源冻结、附件冻结、版本校验和本地 BPM 提交能力。",
    "设计设变、修模改模、标准件、基础资料和 BOM 正式动作已经形成受控适配器。",
    "设计前端已经具备结果表格、图纸预览、参数展示和对话式表单编辑能力。",
    "设计能力已经具备统一权限、证据、审计和人工确认机制。",
])

doc.add_heading("10 依据文件", level=1)
p=doc.add_paragraph()
p.add_run("本报告依据：").bold=True
p.add_run("《模具项目全流程管理系统需求规格说明书 V1.1》、 《模具项目智能工作台技术开发文档 V3.6》、当前 MoldPilot 工作树以及项目内设计工具和需求追踪记录。")
for label, path in [
    ("需求追踪矩阵", r"E:\MoldPilot\docs\REQUIREMENTS_TRACEABILITY.md"),
    ("设计部门工具与 Skill 测试报告", r"E:\MoldPilot\docs\DESIGN_DEPARTMENT_TOOL_SKILL_TEST_REPORT_2026-09-19.md"),
    ("ERP 设计工具与 Skill 清单", r"E:\MoldPilot\docs\ERP_DESIGN_TOOLS_AND_SKILLS.md"),
    ("开发状态记录", r"E:\MoldPilot\docs\DEVELOPMENT_STATUS.md"),
]:
    p=doc.add_paragraph(style="Small Note")
    p.add_run(label+"：").bold=True
    p.add_run(path)

for section in doc.sections:
    footer=section.footer.paragraphs[0]
    footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=footer.add_run("设计部门已完成封装成果报告")
    r.font.name="Microsoft YaHei"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    r.font.size=Pt(8)
    r.font.color.rgb=RGBColor(120,120,120)

doc.save(out)
print(out)
