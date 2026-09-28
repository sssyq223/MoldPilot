# -*- coding: utf-8 -*-
"""Generate a non-technical Word brief on outsource skill progress."""
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

OUT = Path(__file__).resolve().parents[1] / "docs" / "\u59d4\u5916\u5c01\u88c5\u6280\u80fd\u5f00\u53d1\u8fdb\u5ea6\u8bf4\u660e_2026-09-27.docx"

NAVY = RGBColor(0x1F, 0x3A, 0x5F)
ACCENT = RGBColor(0x1A, 0x56, 0x7A)
MUTED = RGBColor(0x4A, 0x55, 0x63)
WARN = RGBColor(0x8A, 0x4B, 0x08)
OK = RGBColor(0x1B, 0x5E, 0x20)
HEADER_FILL = "1A567A"
ALT_FILL = "F4F7FA"
WARN_FILL = "FFF4E5"
OK_FILL = "E8F5E9"


def set_run_font(run, name="微软雅黑", size=11, bold=False, color=None):
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if color is not None:
        run.font.color.rgb = color


def shade_cell(cell, hex_color):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tc_pr.append(shd)


def set_cell_border(cell):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "C5CED8")
        tc_borders.append(el)
    tc_pr.append(tc_borders)


def write_cell(cell, text, *, bold=False, size=10.5, color=None, fill=None, center=False):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color or (RGBColor(0xFF, 0xFF, 0xFF) if fill == HEADER_FILL else RGBColor(0x2B, 0x33, 0x3B)))
    if fill:
        shade_cell(cell, fill)
    set_cell_border(cell)
    cell.vertical_alignment = 1


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16 if level == 1 else 12)
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run(text)
    set_run_font(run, size=16 if level == 1 else 13, bold=True, color=NAVY if level == 1 else ACCENT)
    return p


def add_body(doc, text, *, color=MUTED, space_after=8, first_line=True):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    if first_line:
        p.paragraph_format.first_line_indent = Cm(0.74)
    run = p.add_run(text)
    set_run_font(run, size=11, color=color)
    return p


def add_bullet(doc, text, *, bold_lead=None):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.74)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.35
    if bold_lead:
        run = p.add_run("• " + bold_lead)
        set_run_font(run, size=11, bold=True, color=NAVY)
        run = p.add_run(text)
        set_run_font(run, size=11, color=MUTED)
    else:
        run = p.add_run("• " + text)
        set_run_font(run, size=11, color=MUTED)
    return p


def add_table(doc, headers, rows, col_widths, status_col=None):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for i, width in enumerate(col_widths):
        for cell in table.columns[i].cells:
            cell.width = Cm(width)
    for i, head in enumerate(headers):
        write_cell(table.rows[0].cells[i], head, bold=True, size=10.5, fill=HEADER_FILL, center=True)
    for r, row in enumerate(rows):
        fill = ALT_FILL if r % 2 else "FFFFFF"
        for i, value in enumerate(row):
            color = None
            cell_fill = fill
            if status_col is not None and i == status_col:
                if "尚未" in value or "未测" in value:
                    color = WARN
                    cell_fill = WARN_FILL
                elif "联调" in value or "待再验" in value:
                    color = RGBColor(0x0D, 0x47, 0xA1)
                    cell_fill = "E3F2FD"
                elif "已测" in value or "可用" in value:
                    color = OK
                    cell_fill = OK_FILL
                elif "不做" in value:
                    color = RGBColor(0x5D, 0x40, 0x37)
                    cell_fill = "EFEBE9"
            write_cell(table.rows[r + 1].cells[i], value, size=10, color=color, fill=cell_fill, center=(i != len(row) - 1))
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return table


def main():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.4)
    section.right_margin = Cm(2.4)
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(4)
    run = title.add_run("委外业务智能助手封装")
    set_run_font(run, size=22, bold=True, color=NAVY)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.paragraph_format.space_after = Pt(2)
    run = sub.add_run("开发进度说明")
    set_run_font(run, size=18, bold=True, color=ACCENT)

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.paragraph_format.space_after = Pt(14)
    run = meta.add_run("面向业务同事　·　2026年9月27日　·　内部进度材料")
    set_run_font(run, size=10.5, color=MUTED)

    bar = doc.add_paragraph()
    bar.alignment = WD_ALIGN_PARAGRAPH.CENTER
    bar.paragraph_format.space_after = Pt(12)
    run = bar.add_run("一句话：功能已经按岗位封装好，但真实业务目前只测到「采购填报价」，其余环节还没有走通验收。")
    set_run_font(run, size=11.5, bold=True, color=WARN)

    add_heading(doc, "一、这份材料说什么")
    add_body(
        doc,
        "我们没有另做一套委外系统。现有的零件委外、工序委外、模具委外，仍然以公司现在用的 ERP 为准。"
        "这次做的事情，是把这条链路上各岗位要办的事，封装进统一的对话助手：员工用平常说话的方式提出办理或查询，"
        "助手核对真实单据后弹出确认卡，本人点确认，才写入现有系统。没点确认，不会改数据。",
    )
    add_body(
        doc,
        "下面不写接口、字段、库表这些开发细节，只说明：哪些已经能用对话办理、现在测到哪一步、后面还要补什么。",
    )

    add_heading(doc, "二、当前总进度（请先看这一节）")
    add_body(
        doc,
        "截至今天，六个岗位的查询和办理能力都已经在助手里接好：采购员、采购主管、总经理、加工商、仓管、质检。"
        "员工登录后只能看到自己岗位能办的事，不能代办别人的环节。",
        first_line=True,
    )
    add_body(
        doc,
        "但是，“接好了”不等于“已经用真实业务单验过了”。真实联调目前只覆盖采购员填写我方报价这一步。"
        "发询价曾经碰到过，还不稳定。加工商报价和接单、超区间成交价、主管和总经理审批、仓管发料入库、质检领取和合格，都还没有用真实单据完整走通。",
        first_line=True,
    )

    add_table(
        doc,
        ["进度含义", "现在处在哪"],
        [
            ["功能封装", "各岗位该办的环节，助手里都已接上，按现有 ERP 流程走，不另造一套阶段。"],
            ["真实业务测试", "只测了采购员填我方报价；刚修完口语识别等问题，还要再验一遍。"],
            ["其余办理环节", "代码已接，尚未用真实单据验收，不能当成已经上线可用。"],
            ["本期明确不做", "退换货、对账、拆图合并等异常支线，本轮不纳入。"],
        ],
        [4.2, 12.0],
        status_col=None,
    )

    add_heading(doc, "三、整条委外链，助手打算覆盖什么")
    add_body(doc, "零件委外、模具委外，助手按现有顺序引导，不跳步、不替别人办：", first_line=True)
    add_bullet(doc, "采购员填写我方报价和直接接单上限。")
    add_bullet(doc, "采购员向已匹配的加工商发出询价。")
    add_bullet(doc, "加工商回报价。报价落在区间内，系统按现有规则免审定标；超出区间，采购员再填成交价。")
    add_bullet(doc, "超区间时，采购主管、总经理按节点审批。")
    add_bullet(doc, "加工商接单或拒单。拒单后回到采购重选、再询价。")
    add_bullet(doc, "需要我方供料时，仓管发原料；加工商确认来料、生产、成品发回。")
    add_bullet(doc, "仓管办理回厂到货和入库；质检领取任务并提交合格结论。")
    add_body(
        doc,
        "工序委外更短：没有采购填价、发询价和主管下单审批，按候选加工商一家家发单。接单后仓管备料，加工商发成品，再入库、质检。这两类单子在对话里要分开说，不能把填报价套到工序单上。",
        first_line=True,
    )

    add_heading(doc, "四、各环节封装与测试对照（重点）")
    add_body(doc, "下表是给业务同事看的验收对照。颜色大致表示：蓝色是正在联调，橙色是还没测，绿色是本环节目标状态，灰色是本期不做。", first_line=True)

    add_table(
        doc,
        ["岗位", "助手里能办的事", "封装", "真实测试", "说明"],
        [
            ["采购员", "看自己责任范围内的委外待办", "已接好", "配合填价在测", "查待办、看进度可以问；数量问题应直接出结果，不必先确认一遍。"],
            ["采购员", "填写我方报价、直接接单上限", "已接好", "联调中，待再验", "目前唯一在用真实单据反复测的办理动作。口语说法刚补全，还要再走一遍确认卡。"],
            ["采购员", "向已匹配加工商发询价", "已接好", "尚未稳定验收", "碰过确认卡出不来、加工商名称被说错等问题，不能算过关。"],
            ["采购员", "超区间填写成交价、提交审批", "已接好", "尚未测试", "要等前面询价、加工商回价走完，才有真实单可测。"],
            ["采购员", "拒单后重选加工商", "已接好（说明下一步）", "尚未测试", "现有系统没有单独的重选按钮，助手只说明下一步应再发询价，不直接改系统。"],
            ["采购主管 / 总经理", "超区间下单审批：通过或驳回", "已接好", "尚未测试", "两人共用同一套办理，按各自审批节点看自己的待办。"],
            ["加工商", "看本厂待办、提交报价", "已接好", "尚未测试", "只能看自己的单，看不到别家和内部价。"],
            ["加工商", "接单 / 拒单", "已接好", "尚未测试", "工序拒单后由现有系统转下一家，加工商不用自己选下一家。"],
            ["加工商", "确认来料、成品发回", "已接好", "尚未测试", "没收到的料不能按整单数量发成品。"],
            ["仓管", "原料发出 / 工序备料完成", "已接好", "尚未测试", "采购直发不是仓库待办，助手不会拿来给仓管办。"],
            ["仓管", "回厂到货确认、入库", "已接好", "尚未测试", "可以按行、按数量部分收；拒收本轮不办。"],
            ["质检", "领取任务、提交全检合格", "已接好", "尚未测试", "不合格、退货本轮不办。"],
            ["各岗位", "退换货、对账、拆图合并", "不在本轮", "本期不做", "现有系统里有的异常支线，本轮助手不封装。"],
        ],
        [3.0, 4.4, 2.2, 2.8, 3.8],
        status_col=3,
    )

    add_heading(doc, "五、采购填报价：现在测到什么程度")
    add_body(
        doc,
        "采购员在对话里说出模具、批次或零件，并给出总价和上限后，助手应直接弹出报价确认卡。"
        "卡片上核对订单（尚未下单就写尚未下单）、模具、批次、零件和金额。本人点确认后，才写入现有系统。没有确认卡、只口头复述金额，不能算办成。",
        first_line=True,
    )
    add_body(doc, "联调过程中已经碰到、并正在收口的问题：", first_line=True)
    add_bullet(doc, "员工说「报报价」时，助手原先只认「填我方报价」这类说法，会回成「还没有查」。口语已经补上，需要再用真实单再试一次。", bold_lead="说话认不全。")
    add_bullet(doc, "零件号和中文名称连在一起时，有时锁不到正确的那一张询价。识别已经收紧，仍要再验「待排第一张」这种说法。", bold_lead="单子对不齐。")
    add_bullet(doc, "确认卡有时出不来，或点确认后提示不像业务结果。这类问题修过一轮，填价确认还要再点一遍看是否干净。", bold_lead="确认卡不稳。")
    add_bullet(doc, "发询价时，助手曾把已匹配的加工商说成另一个近似名称。规则已改为：没点名就用系统里已匹配的名单，禁止编名字。发询价本身还没验收。", bold_lead="加工商名称说错。")
    add_body(
        doc,
        "所以：采购填报价是目前唯一进入真实联调的办理环节，但还不能对外说「已经测完、可以放手用」。"
        "要等确认卡能稳定弹出、点确认后现有系统里的金额正确，这一步才算过关。",
        first_line=True,
    )

    add_heading(doc, "六、还没测到、但后面必须补上的")
    add_body(doc, "建议严格按现有流程往下走，前面没验稳不要跳到后面，否则真实单到不了下一岗。", first_line=True)
    add_bullet(doc, "采购员对已匹配加工商发询价，确认卡上的厂家必须和现有系统里的名单一致。", bold_lead="下一步优先：")
    add_bullet(doc, "加工商回报价；区间内是否自动到待接单，超区间是否正确落到采购填成交价。")
    add_bullet(doc, "采购员填成交价，主管、总经理按节点审批。")
    add_bullet(doc, "加工商接单、拒单；零件/模具拒单后采购重选，工序拒单后是否自动转下一家。")
    add_bullet(doc, "仓管发料或备料、回厂入库；加工商确认来料和成品发回。")
    add_bullet(doc, "质检领取和合格。不合格、退货仍不在本轮。")
    add_bullet(doc, "工序委外整条短链，要单独用工序单测，不能拿零件填价的结论套过去。")
    add_body(
        doc,
        "测试时还要注意：同一模具常有多张询价，必须点到零件或「第几张」；助手不能猜。"
        "换账号就要换岗位，采购员账号办不了接单，加工商账号办不了填价。",
        first_line=True,
    )

    add_heading(doc, "七、给业务同事的结论")
    add_table(
        doc,
        ["可以怎么理解", "现在还不能怎么理解"],
        [
            ["助手已经按岗位把委外办理接进对话，方向对、范围清楚。", "不能理解成整条委外链已经用真实单验收完毕。"],
            ["采购填报价正在联调，问题在收口，很快可以再验。", "不能理解成采购员所有动作（含发询价、成交价）都已可用。"],
            ["后面各岗的办理卡片已经按现有流程备好，具备开测条件。", "不能理解成加工商、仓管、质检、审批现在就能放手给一线用。"],
            ["没点确认卡，不会改现有系统。", "不能理解成「助手说办成了」就等于系统里已经办成。以确认卡和系统回执为准。"],
        ],
        [8.1, 8.1],
    )

    add_heading(doc, "八、建议的下一步")
    add_body(doc, "按下面顺序推进，每过一关再开下一关：", first_line=True)
    add_bullet(doc, "用真实待填价单据，把「报报价 / 待排第一张 / 点名模具和零件」三条说法再走一遍，确认卡弹出、金额写入现有系统。", bold_lead="本周先收口：")
    add_bullet(doc, "同一张单接着发询价，核对加工商名单，本人确认后看系统是否真正发出。")
    add_bullet(doc, "换加工商账号走报价、接单；有超区间单再走成交价和两级审批。")
    add_bullet(doc, "仓管、质检用履约中的真实单各走一遍。")
    add_bullet(doc, "另备一张工序委外单，单独走短链，避免和零件填价混测。")

    note = doc.add_paragraph()
    note.paragraph_format.space_before = Pt(18)
    note.paragraph_format.space_after = Pt(4)
    run = note.add_run("说明")
    set_run_font(run, size=10, bold=True, color=ACCENT)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(
        "本材料只反映截至 2026年9月27日 的封装与联调进度，供内部沟通。"
        "权威数据与能否办理，一律以现有 ERP 当时状态为准。助手是办理入口，不是第二套台账。"
    )
    set_run_font(run, size=9.5, color=MUTED)

    footer = doc.add_paragraph()
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = footer.add_run("MoldPilot　委外封装　内部进度")
    set_run_font(run, size=9, color=RGBColor(0x90, 0xA0, 0xB0))

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
