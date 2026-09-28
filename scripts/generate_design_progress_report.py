
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_BREAK
from pathlib import Path

out = Path(r"E:\设计部门封装进度分析报告.docx")

doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(0.7)
sec.bottom_margin = Inches(0.65)
sec.left_margin = Inches(0.75)
sec.right_margin = Inches(0.75)

# Base styles
styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Microsoft YaHei"
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
normal.font.size = Pt(10.5)
normal.paragraph_format.space_after = Pt(5)
normal.paragraph_format.line_spacing = 1.18

for name, size, color in [("Title", 22, "000000"), ("Heading 1", 15, "000000"), ("Heading 2", 12.5, "000000"), ("Heading 3", 11, "000000")]:
    st = styles[name]
    st.font.name = "Microsoft YaHei"
    st._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = RGBColor.from_string(color)
    st.paragraph_format.space_before = Pt(10 if name != "Title" else 0)
    st.paragraph_format.space_after = Pt(5)

# custom styles
if "Report Subtitle" not in styles:
    st = styles.add_style("Report Subtitle", WD_STYLE_TYPE.PARAGRAPH)
    st.font.name = "Microsoft YaHei"
    st._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    st.font.size = Pt(12)
    st.font.color.rgb = RGBColor(80,80,80)
    st.paragraph_format.space_after = Pt(12)
if "Small Note" not in styles:
    st = styles.add_style("Small Note", WD_STYLE_TYPE.PARAGRAPH)
    st.font.name = "Microsoft YaHei"
    st._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    st.font.size = Pt(9)
    st.font.color.rgb = RGBColor(95,95,95)
    st.paragraph_format.space_after = Pt(4)

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcPr.append(shd)
    shd.set(qn("w:fill"), fill)

def set_cell_border(cell, color="D9D9D9", sz="4"):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    borders = tcPr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tcPr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        el = borders.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), sz)
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)

def set_cell_margin(cell, top=90, start=100, bottom=90, end=100):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tcMar.find(qn("w:" + m))
        if node is None:
            node = OxmlElement("w:" + m)
            tcMar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")

def set_repeat_table_header(row):
    trPr = row._tr.get_or_add_trPr()
    tblHeader = OxmlElement("w:tblHeader")
    tblHeader.set(qn("w:val"), "true")
    trPr.append(tblHeader)

def write_cell(cell, text, bold=False, color=None, align=None, size=9.2):
    cell.text = ""
    p = cell.paragraphs[0]
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(str(text))
    r.bold = bold
    r.font.name = "Microsoft YaHei"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    r.font.size = Pt(size)
    if color:
        r.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margin(cell)
    set_cell_border(cell)

def add_table(headers, rows, widths=None, header_fill="3F4E5E"):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for i, h in enumerate(headers):
        write_cell(hdr.cells[i], h, bold=True, color="FFFFFF", align=WD_ALIGN_PARAGRAPH.CENTER, size=9.3)
        shade(hdr.cells[i], header_fill)
    for ridx, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            align = WD_ALIGN_PARAGRAPH.CENTER if i == 0 else WD_ALIGN_PARAGRAPH.LEFT
            write_cell(cells[i], value, align=align, size=9.0)
            if ridx % 2 == 1:
                shade(cells[i], "F5F7F9")
        for c in cells:
            set_cell_border(c)
    if widths:
        for row in table.rows:
            for i, width in enumerate(widths):
                row.cells[i].width = Inches(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table

def add_bullets(items, level=0):
    for item in items:
        p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
        p.paragraph_format.space_after = Pt(2)
        p.add_run(item)

def add_numbered(items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(2)
        p.add_run(item)

# Title
p = doc.add_paragraph(style="Title")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run("设计部门封装进度分析报告")
p = doc.add_paragraph(style="Report Subtitle")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run("基于需求规格 V1.1、技术开发文档 V3.6 与当前 MoldPilot 工作树")
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("2026年9月27日")
r.font.size = Pt(10)
r.font.color.rgb = RGBColor(100,100,100)

doc.add_paragraph()
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(8)
r = p.add_run("结论摘要：")
r.bold = True
p.add_run("设计部门的工具和 Skill 封装已经基本成型，代码封装约 80%～85%；但正式设计版本、设计委外排程、设变闭环、试模料标准和真实 ERP 联调尚未完成，可交付和可验收进度约 45%～50%。")

add_table(
    ["统计口径", "当前结果", "说明"],
    [
        ["设计 Skill", "22 个", "包含设计路线、上传、BOM、设变、修模、标准件、订单等能力"],
        ["设计相关工具", "60 个", "58 个 ERP 设计工具，1 个本地设计路线查询，1 个设计审批准备工具"],
        ["非正式写入工具", "33 个", "查询、解析、校验、核价、公差、图纸、分析和下载等"],
        ["正式动作工具", "25 个", "当前工作树已接入人工确认卡，确认后才调用 ERP 控制接口"],
        ["需求验收状态", "FR-043～FR-047 均未验证", "已有代码证据，但缺真实业务闭环和 ERP 端到端证据"],
    ],
    widths=[1.35, 1.35, 4.6]
)

doc.add_heading("1 评估范围与判定方法", level=1)
doc.add_paragraph("本报告重点分析设计部门封装，不把 ERP 原有功能重复计算为 MoldPilot 自主开发。凡是调用 management-system ERP 的能力，按“适配器已封装、人工确认边界、真实接口已联调、业务闭环已验收”分别判断。")
doc.add_paragraph("当前代码已经存在大量工具和页面展示，因此“代码封装进度”高于“可交付进度”。需求追踪矩阵仍将 FR-043～FR-047 标记为 NOT_VERIFIED，说明这些条目不能仅凭单元测试或工具注册视为完成。", style="Small Note")

doc.add_heading("2 已封装能力总览", level=1)
add_table(
    ["能力域", "主要已封装内容", "当前进度", "状态"],
    [
        ["设计路线与项目上下文", "设计版本、BOM、路线、计划任务、工程联络影响查询与改版差异分析", "约70%", "部分完成"],
        ["新模/改模清单", "文件解析、会话、状态、结果、校验、核价、公差、参数、图纸预览和导入适配", "约80%", "基本成型"],
        ["设计订单审批材料", "ERP 订单快照、版本、哈希、附件冻结、本地 Agent BPM 提交", "约65%～70%", "部分完成"],
        ["BOM 与工艺路线", "BOM 查询、报表、缺料、导入、维护及本地路线聚合", "约60%", "部分完成"],
        ["设变与修模", "版本对比、设变分析、修模审批读取和正式动作适配", "约55%", "部分完成"],
        ["基础资料与标准件", "密度、分组规则、关键词、标准件目录查询及维护动作适配", "查询约80%，维护约50%", "查询较完整"],
        ["统一工作台前端", "证据卡片、表格、图纸预览、参数编辑、上传结果呈现", "约80%", "基本成型"],
        ["权限与确认控制", "能力目录、权限范围、人工确认卡、版本和幂等校验", "约70%", "正在收口"],
    ],
    widths=[1.35, 3.35, 1.15, 1.15]
)

doc.add_heading("3 已完成或基本完成的封装内容", level=1)

doc.add_heading("3.1 设计路线和改版影响查询", level=2)
add_bullets([
    "提供 query_design_route_context，按项目、设计单、图纸版本、BOM 物料、计划任务或工程联络线索定位对象。",
    "返回设计版本、BOM 明细、内部加工/采购/委外路线、关联计划任务和工程联络影响。",
    "支持上一版与当前版的 BOM 新增、删除、数量变化、路线变化和任务关联变化分析。",
    "能够生成计划变更复核候选，但不会直接修改计划、顺延客户交期或下达采购/加工任务。",
])

doc.add_heading("3.2 新模和改模清单流程", level=2)
add_bullets([
    "新模和改模分别使用独立解析入口，改模流程固定使用 repair_other 业务类型。",
    "已经封装上传会话、图纸处理状态、解析结果、明细校验和审批配置查询。",
    "已经封装 ERP 钢料价格重算、公差计算、材质/数量/尺寸参数查询和图纸预览。",
    "已经封装按图纸自动修正结果的前后差异展示，并在前端支持分页、筛选和字段编辑。",
    "新模导入、改模导入、变更请购等正式动作已经登记为 ERP 适配器。",
])

doc.add_heading("3.3 设计审批材料和本地 BPM", level=2)
add_bullets([
    "prepare_design_order_approval 会先读取 ERP 设计订单，再形成 MoldPilot 确认卡。",
    "审批材料冻结 ERP 来源、资源 ID、版本、核对时间、快照哈希、订单摘要和本轮附件版本。",
    "确认时重新读取 ERP，订单发生变化会返回 VERSION_CONFLICT，阻止使用过期材料。",
    "确认成功后创建本地 design_route 并提交 Agent BPM，不直接调用 ERP 旧审批流程。",
])

doc.add_heading("3.4 ERP 设计工作台查询", level=2)
add_bullets([
    "设计订单、图纸版本、BOM、BOM 报表、缺料、设变、标准件、闲置料、材质密度、分组规则和关键词均有查询工具。",
    "支持单记录读取、图纸版本对比和设变影响分析。",
    "修模改模相关的普通审批、委外审批和加工商响应均有只读查询入口。",
])

doc.add_heading("3.5 正式动作适配器", level=2)
add_bullets([
    "设计订单明细修改、闲置料保存/释放、订单生命周期、BOM 维护、设变维护、设变明细、标准件维护、修模处理和清单导入均已登记。",
    "当前工作树已将 ERP 设计正式动作统一转入确认卡：模型首次调用生成提案，人工确认接口再执行 ERP 控制工具。",
    "确认流程具备用户、权限版本、提案哈希和确认凭证校验，并保留执行回执。",
])

doc.add_heading("4 各需求条目进度", level=1)
add_table(
    ["需求", "已经做到", "仍缺少", "进度判断"],
    [
        ["FR-043 内部设计/设计委外", "有设计路线查询、设计审批材料和本地 BPM 提交", "设计主管确认、负责人/费用/供应商、设计排期、委外成果接收与整改", "约60%，未验收"],
        ["FR-044 工艺分析到设计确认", "有上传解析、图纸、参数、公差和设计上下文查询", "工艺分析任务、结构设计、出图、设计确认、委外审核和整改", "约40%，未验收"],
        ["FR-045 正式版本/BOM/路线/任务", "有本地设计路线、ERP BOM 查询/导入/维护适配", "正式设计版本发布、统一生效、下游任务正式绑定和联动", "约40%，未验收"],
        ["FR-046 改版和工程联络转设计", "有图纸版本对比、改版影响分析、计划复核候选、设变工具", "工程联络转设计订单、批准后版本执行、采购/加工/供应商任务闭环", "约55%，未验收"],
        ["FR-047 试模料标准", "有技术要求、公差和清单数据读取", "客户/规格/颜色、发货地点、招标、无合同确认记录和附件来源", "约25%，未验收"],
    ],
    widths=[1.35, 2.35, 2.75, 1.2]
)

doc.add_heading("5 尚未完成的主要内容", level=1)
doc.add_heading("5.1 正式设计成果发布", level=2)
add_bullets([
    "尚未形成完整的正式设计版本发布流程。",
    "尚未把设计确认后的图纸、BOM、工艺路线、零件清单统一冻结为执行依据。",
    "尚未完成正式版本与采购、加工、装配、试模任务的完整绑定。",
])

doc.add_heading("5.2 设计排产和设计委外", level=2)
add_bullets([
    "设计排产目前仍以线下安排和线上进度记录为主。",
    "缺少设计委外任务下达、供应商成果接收、审核、退回整改和重新提交的完整链路。",
    "费用、供应商、计划日期和设计负责人记录还没有形成完整可验收业务闭环。",
])

doc.add_heading("5.3 设变和工程联络闭环", level=2)
add_bullets([
    "目前已经能查询和分析影响，但不能据此证明设变已经落实。",
    "工程联络单转设计订单、影响对象执行、整改、复验和关闭仍需继续开发。",
    "设计改版对采购、加工、外协和交期的联动仍以候选和建议为主。",
])

doc.add_heading("5.4 试模料标准", level=2)
add_bullets([
    "尚未建立完整的试模料标准业务记录。",
    "客户供料判断、额外采购数量、发货地点和无合同确认依据还没有完整结构化。",
    "设计与项目并行确认以及采购交接证据还没有完成。",
])

doc.add_heading("5.5 真实 ERP 联调", level=2)
add_bullets([
    "当前没有真实 ERP 登录、真实客户文件和真实生产接口回执作为验收证据。",
    "用户指定的 D:\\ERP-system\\management-system 路径在当前环境不可用；项目已有记录显示曾发现其他 ERP 工作区，但本轮没有连接生产 ERP。",
    "ERP 适配器的接口参数、权限、重复提交、版本冲突和异常回执需要在 ERP 启动后统一验证。",
])

doc.add_heading("6 当前验证证据", level=1)
add_table(
    ["验证项", "结果", "说明"],
    [
        ["设计能力目录定向测试", "13 项通过", "覆盖设计工具、Skill、确认模式和能力元数据"],
        ["前端设计预览测试", "31 项通过", "覆盖上传预览、参数、公差、图纸和表单命令"],
        ["Vue 类型检查", "通过", "mold 前端类型检查通过"],
        ["前端生产构建", "通过", "生产包构建成功"],
        ["ERP 设计适配器测试", "28 项通过，4 项失败", "部分旧测试仍期望首次调用直接执行 ERP 写入"],
        ["设计审批数据库测试", "部分未执行", "测试配置指向 192.168.3.215，被安全守卫阻止"],
        ["真实 ERP 端到端", "未完成", "没有真实 ERP 登录和业务回执证据"],
    ],
    widths=[1.7, 1.7, 4.2]
)

doc.add_heading("7 当前工作树中的关键风险", level=1)
add_bullets([
    "写操作确认机制正在从“模型传 confirm=true”迁移到“人工确认卡 + 一次性确认凭证”，实现已经存在，但适配器测试和完整回归尚未全部更新。",
    "部分自然语言仍可能触发过宽或错误的 Skill，例如自动修正、修模 DXF 上传和 BOM XLSX 导入场景，需要继续收窄路由。",
    "设计工具清单文档与当前代码数量存在漂移，文档仍按旧版本统计，需统一到当前 22 个 Skill、60 个设计工具。",
    "工具有代码不等于真实业务已完成；当前所有 ERP 业务动作仍应以真实 ERP 回执、人工审批记录和完整追溯证据为准。",
])

doc.add_heading("8 后续开发优先级", level=1)
add_numbered([
    "先完成 25 个正式动作的确认卡测试迁移、一次性凭证、重复确认和异常回执回归。",
    "修复设计 Skill 路由冲突，确保自动修正、修模 DXF 上传和 BOM XLSX 导入进入正确流程。",
    "实现正式设计版本发布，并建立图纸、BOM、工艺路线和下游任务的统一版本关联。",
    "补齐工程联络单转设计订单、设变执行、整改、复验和关闭。",
    "补齐试模料标准、客户供料判断和设计与项目并行确认。",
    "ERP 启动后统一联调上传、导入、订单、BOM、设变、修模、标准件和版本冲突场景。",
    "最后同步工具清单、追踪矩阵和开发状态文档，避免统计口径继续漂移。",
])

doc.add_heading("9 依据文件", level=1)
p = doc.add_paragraph()
p.add_run("本报告依据：").bold = True
p.add_run("《模具项目全流程管理系统需求规格说明书 V1.1》、 《模具项目智能工作台技术开发文档 V3.6》、当前 MoldPilot 工作树以及项目内的需求追踪和设计专项测试记录。")
for path, label in [
    (r"E:\MoldPilot\docs\REQUIREMENTS_TRACEABILITY.md", "需求追踪矩阵"),
    (r"E:\MoldPilot\docs\DESIGN_DEPARTMENT_TOOL_SKILL_TEST_REPORT_2026-09-19.md", "设计部门工具与 Skill 测试报告"),
    (r"E:\MoldPilot\docs\ERP_DESIGN_TOOLS_AND_SKILLS.md", "ERP 设计工具与 Skill 清单"),
    (r"E:\MoldPilot\docs\DEVELOPMENT_STATUS.md", "开发状态记录"),
]:
    p = doc.add_paragraph(style="Small Note")
    p.add_run(label + "：").bold = True
    p.add_run(path)

# Footer
for section in doc.sections:
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("设计部门封装进度分析报告")
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(120,120,120)

doc.save(out)
print(out)
