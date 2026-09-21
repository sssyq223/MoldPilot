"""High-fidelity office document previews for the local workbench."""
from __future__ import annotations

import os
import csv
import io
import shutil
import subprocess
import threading
from hashlib import sha256
from pathlib import Path
from tempfile import TemporaryDirectory
from datetime import date, datetime, time

from .errors import DomainError
from .config import settings


_lock_guard = threading.Lock()
_preview_locks: dict[str, threading.Lock] = {}

# Spreadsheet previews are deliberately bounded.  The original file remains
# available through the separately authorised download endpoint; the preview
# endpoint only returns a small, plain-value table and never evaluates macros,
# formulas, links or embedded workbook objects.
SPREADSHEET_PREVIEW_MAX_ROWS = 200
SPREADSHEET_PREVIEW_MAX_COLUMNS = 50
SPREADSHEET_PREVIEW_MAX_CELL_LENGTH = 2000
SPREADSHEET_PREVIEW_MAX_SHEETS = 20


_WORD_EXPORT_SCRIPT = r"""
param([string]$InputPath, [string]$OutputPath)
$ErrorActionPreference = 'Stop'
$word = $null
$documents = $null
$document = $null
$options = $null
try {
    $word = New-Object -ComObject Word.Application
    $word.Visible = $false
    $word.DisplayAlerts = 0
    $word.AutomationSecurity = 3
    $options = $word.Options
    $options.UpdateLinksAtOpen = $false
    $documents = $word.Documents
    $document = $documents.Open($InputPath, $false, $true)
    $document.ExportAsFixedFormat($OutputPath, 17)
}
finally {
    if ($null -ne $document) {
        $document.Close($false)
        [void][Runtime.InteropServices.Marshal]::FinalReleaseComObject($document)
    }
    if ($null -ne $documents) {
        [void][Runtime.InteropServices.Marshal]::FinalReleaseComObject($documents)
    }
    if ($null -ne $options) {
        [void][Runtime.InteropServices.Marshal]::FinalReleaseComObject($options)
    }
    if ($null -ne $word) {
        $word.Quit()
        [void][Runtime.InteropServices.Marshal]::FinalReleaseComObject($word)
    }
    [GC]::Collect()
    [GC]::WaitForPendingFinalizers()
}
"""


def _run(command: list[str], timeout: int = 60) -> None:
    flags = subprocess.CREATE_NO_WINDOW if os.name == "nt" else 0
    result = subprocess.run(command, capture_output=True, text=True, timeout=timeout,
                            check=False, creationflags=flags)
    if result.returncode:
        detail = (result.stderr or result.stdout or "document converter failed").strip()
        raise RuntimeError(detail[-1000:])


def _word_to_pdf(source: Path, target: Path, work: Path) -> None:
    script = work / "export-word-preview.ps1"
    script.write_text(_WORD_EXPORT_SCRIPT, encoding="utf-8")
    powershell = shutil.which("powershell.exe") or shutil.which("powershell")
    if not powershell:
        raise RuntimeError("Microsoft Word preview converter is unavailable")
    _run([powershell, "-NoProfile", "-NonInteractive", "-ExecutionPolicy", "Bypass",
          "-File", str(script), str(source), str(target)])


def _libreoffice_to_pdf(source: Path, target: Path, work: Path) -> None:
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        raise RuntimeError("LibreOffice preview converter is unavailable")
    _run([soffice, "--headless", "--convert-to", "pdf", "--outdir", str(work), str(source)])
    generated = work / (source.stem + ".pdf")
    if generated != target and generated.exists():
        generated.replace(target)


def _lock_for(digest: str) -> threading.Lock:
    with _lock_guard:
        return _preview_locks.setdefault(digest, threading.Lock())


def _cached_pdf(digest: str) -> Path:
    root = Path(settings().file_local_root).resolve().parent / "preview-cache"
    return root / f"{digest}.pdf"


def docx_to_pdf(data: bytes, digest: str | None = None) -> bytes:
    """Render a DOCX through the installed office engine, preserving Word layout."""
    digest = digest or sha256(data).hexdigest()
    cache = _cached_pdf(digest)
    try:
        with _lock_for(digest):
            if cache.exists():
                rendered = cache.read_bytes()
                if rendered.startswith(b"%PDF-"):
                    return rendered
            with TemporaryDirectory(prefix="moldpilot-docx-preview-") as directory:
                work = Path(directory)
                source, target = work / "source.docx", work / "preview.pdf"
                source.write_bytes(data)
                if os.name == "nt":
                    _word_to_pdf(source, target, work)
                else:
                    _libreoffice_to_pdf(source, target, work)
                rendered = target.read_bytes()
                if not rendered.startswith(b"%PDF-"):
                    raise RuntimeError("converter did not produce a valid PDF")
                cache.parent.mkdir(parents=True, exist_ok=True)
                temporary = cache.with_suffix(f".{os.getpid()}.tmp")
                temporary.write_bytes(rendered);temporary.replace(cache)
                return rendered
    except (OSError, subprocess.SubprocessError, RuntimeError):
        raise DomainError("PREVIEW_RENDER_FAILED", "Word 文档暂时无法高保真预览，请下载原件查看", 503)


def _preview_value(value: object) -> str | int | float | bool | None:
    """Convert spreadsheet values to JSON-safe display values.

    Dates are rendered in ISO form and everything else is kept as a primitive
    where possible.  This is a display-only projection, not an import parser.
    """
    if value is None:
        return None
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, (str, int, float, bool)):
        if isinstance(value, str) and len(value) > SPREADSHEET_PREVIEW_MAX_CELL_LENGTH:
            return value[:SPREADSHEET_PREVIEW_MAX_CELL_LENGTH] + "…"
        return value
    text = str(value)
    suffix = "…" if len(text) > SPREADSHEET_PREVIEW_MAX_CELL_LENGTH else ""
    return text[:SPREADSHEET_PREVIEW_MAX_CELL_LENGTH] + suffix


def _excel_color(color: object) -> str | None:
    """Return a CSS color for the RGB colors exposed by openpyxl."""
    if color is None or getattr(color, "type", None) != "rgb":
        return None
    value = str(getattr(color, "rgb", "") or "")
    if len(value) == 8:
        value = value[-6:]
    if len(value) != 6 or value.lower() == "000000":
        return None
    return f"#{value}"


def _border_side(side: object) -> dict[str, str] | None:
    style = str(getattr(side, "style", "") or "")
    if not style:
        return None
    result = {"style": style}
    color = _excel_color(getattr(side, "color", None))
    if color:
        result["color"] = color
    return result


def _xlsx_cell_style(cell: object) -> dict[str, object]:
    """Project the safe, display-only portion of an XLSX cell style."""
    font = getattr(cell, "font", None)
    fill = getattr(cell, "fill", None)
    alignment = getattr(cell, "alignment", None)
    border = getattr(cell, "border", None)
    style: dict[str, object] = {}
    if font:
        font_style: dict[str, object] = {
            "name": str(getattr(font, "name", "") or ""),
            "size": float(getattr(font, "sz", 11) or 11),
            "bold": bool(getattr(font, "b", False)),
            "italic": bool(getattr(font, "i", False)),
        }
        color = _excel_color(getattr(font, "color", None))
        if color:
            font_style["color"] = color
        style["font"] = font_style
    if fill and str(getattr(fill, "fill_type", "") or "") == "solid":
        color = _excel_color(getattr(fill, "fgColor", None))
        if color:
            style["fill"] = color
    if alignment:
        style["alignment"] = {
            "horizontal": getattr(alignment, "horizontal", None),
            "vertical": getattr(alignment, "vertical", None),
            "wrap": bool(getattr(alignment, "wrap_text", False)),
            "rotation": int(getattr(alignment, "text_rotation", 0) or 0),
            "indent": int(getattr(alignment, "indent", 0) or 0),
        }
    if border:
        style["border"] = {
            side: value for side, value in (
                ("left", _border_side(getattr(border, "left", None))),
                ("right", _border_side(getattr(border, "right", None))),
                ("top", _border_side(getattr(border, "top", None))),
                ("bottom", _border_side(getattr(border, "bottom", None))),
            ) if value
        }
    number_format = str(getattr(cell, "number_format", "") or "")
    if number_format and number_format != "General":
        style["number_format"] = number_format
    return style


def _xlsx_rich_grid(sheet: object, max_row: int, max_column: int) -> list[list[dict[str, object]]]:
    """Build a bounded grid retaining merges and display styles for the UI."""
    merged: dict[tuple[int, int], tuple[int, int, int, int]] = {}
    covered: set[tuple[int, int]] = set()
    for area in getattr(sheet, "merged_cells", ()).ranges:
        # openpyxl exposes bounds as (min_col, min_row, max_col, max_row).
        min_col, min_row, max_col_area, max_row_area = area.bounds
        min_row = max(1, min_row); min_col = max(1, min_col)
        max_row_area = min(max_row, max_row_area); max_col_area = min(max_column, max_col_area)
        if min_row > max_row_area or min_col > max_col_area:
            continue
        span = (min_row, min_col, max_row_area, max_col_area)
        merged[(min_row, min_col)] = span
        for row in range(min_row, max_row_area + 1):
            for column in range(min_col, max_col_area + 1):
                if (row, column) != (min_row, min_col):
                    covered.add((row, column))
    result: list[list[dict[str, object]]] = []
    for row in range(1, max_row + 1):
        rendered: list[dict[str, object]] = []
        column = 1
        while column <= max_column:
            if (row, column) in covered:
                column += 1
                continue
            cell = sheet.cell(row, column)
            item: dict[str, object] = {
                "column": column - 1,
                "value": _preview_value(cell.value),
                "style": _xlsx_cell_style(cell),
            }
            area = merged.get((row, column))
            if area:
                item["row_span"] = area[2] - area[0] + 1
                item["column_span"] = area[3] - area[1] + 1
                column = area[3] + 1
            else:
                column += 1
            rendered.append(item)
        result.append(rendered)
    return result


def _bounded_rows(rows) -> tuple[list[list[object]], int, int, bool]:
    """Normalize an iterable of rows and return rows/counts/truncation state."""
    output: list[list[object]] = []
    max_columns = 0
    truncated = False
    source_count = 0
    for source_count, row in enumerate(rows, 1):
        # Stop after one sentinel row.  Counting every row in a hostile or
        # very large workbook would defeat the resource bound this preview
        # parser is intended to provide.
        if source_count > SPREADSHEET_PREVIEW_MAX_ROWS:
            truncated = True
            break
        values: list[object] = []
        for column_index, value in enumerate(row or ()):
            if column_index >= SPREADSHEET_PREVIEW_MAX_COLUMNS:
                truncated = True
                break
            values.append(value)
        normalized = [_preview_value(value) for value in values]
        max_columns = max(max_columns, len(normalized))
        output.append(normalized)
    # ``row_count`` is the number actually returned.  When truncated, the
    # explicit flag tells the UI that additional rows exist without forcing a
    # full scan just to compute an exact total.
    return output, (len(output) if truncated else source_count), max_columns, truncated


def _csv_preview(data: bytes) -> list[dict[str, object]]:
    # BOM-aware UTF-8 first, then the common Chinese office encodings.  Decode
    # errors are replaced so a damaged character cannot break the whole file
    # preview; the original bytes are still downloadable for exact recovery.
    text = ""
    for encoding in ("utf-8-sig", "utf-16", "gb18030", "big5"):
        try:
            text = data.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    if not text:
        text = data.decode("utf-8", errors="replace")
    try:
        dialect = csv.Sniffer().sniff(text[:8192], delimiters=",\t;|")
    except csv.Error:
        dialect = csv.excel
    reader = csv.reader(io.StringIO(text, newline=""), dialect=dialect)
    rows, row_count, column_count, truncated = _bounded_rows(reader)
    column_widths = [
        min(40.0, max(8.43, max((len(str(row[index])) for row in rows if index < len(row)), default=8.43)))
        for index in range(column_count)
    ]
    return [{
        "name": "CSV",
        "rows": rows,
        "row_count": row_count,
        "column_count": column_count,
        "truncated": truncated,
        "column_widths": column_widths,
    }]


def _xlsx_preview(data: bytes) -> list[dict[str, object]]:
    try:
        from openpyxl import load_workbook
        from openpyxl.utils import get_column_letter

        # A normal workbook is required here because read-only mode discards
        # the dimensions, merged ranges and cell styles needed for a faithful
        # display.  The preview remains bounded and never evaluates formulas.
        workbook = load_workbook(filename=io.BytesIO(data), read_only=False, data_only=True, keep_links=False)
        sheets: list[dict[str, object]] = []
        try:
            for sheet in list(workbook.worksheets)[:SPREADSHEET_PREVIEW_MAX_SHEETS]:
                max_row = min(int(sheet.max_row or 1), SPREADSHEET_PREVIEW_MAX_ROWS)
                max_column = min(int(sheet.max_column or 1), SPREADSHEET_PREVIEW_MAX_COLUMNS)
                rows, row_count, column_count, truncated = _bounded_rows(
                    sheet.iter_rows(values_only=True, max_row=max_row, max_col=max_column)
                )
                # Empty worksheet tabs are implementation leftovers in many
                # Excel templates; showing them adds no previewable content.
                if not any(value is not None for row in rows for value in row):
                    continue
                column_widths = [
                    float(sheet.column_dimensions[get_column_letter(column)].width or 8.43)
                    for column in range(1, max_column + 1)
                ]
                row_heights = [
                    float(sheet.row_dimensions[row].height)
                    if sheet.row_dimensions[row].height is not None else None
                    for row in range(1, max_row + 1)
                ]
                sheets.append({
                    "name": str(sheet.title),
                    "rows": rows,
                    "row_count": max(row_count, max_row),
                    "column_count": max(column_count, max_column),
                    "truncated": truncated or int(sheet.max_row or 1) > max_row or int(sheet.max_column or 1) > max_column or len(workbook.worksheets) > SPREADSHEET_PREVIEW_MAX_SHEETS,
                    "column_widths": column_widths,
                    "row_heights": row_heights,
                    "grid": _xlsx_rich_grid(sheet, max_row, max_column),
                })
        finally:
            workbook.close()
        return sheets
    except Exception as error:  # parser errors are normalized for the API boundary
        raise DomainError("PREVIEW_RENDER_FAILED", "Excel 文件无法解析为安全预览，请下载原件查看", 503) from error


def _xls_preview(data: bytes) -> list[dict[str, object]]:
    try:
        import xlrd

        workbook = xlrd.open_workbook(file_contents=data, on_demand=True, ragged_rows=True)
        sheets: list[dict[str, object]] = []
        for sheet in workbook.sheets()[:SPREADSHEET_PREVIEW_MAX_SHEETS]:
            def rows():
                for row_index in range(sheet.nrows):
                    limit = min(sheet.row_len(row_index), SPREADSHEET_PREVIEW_MAX_COLUMNS + 1)
                    yield [sheet.cell_value(row_index, col) for col in range(limit)]

            preview_rows, row_count, column_count, truncated = _bounded_rows(rows())
            if not any(value is not None for row in preview_rows for value in row):
                continue
            sheets.append({
                "name": str(sheet.name),
                "rows": preview_rows,
                "row_count": row_count,
                "column_count": column_count,
                "truncated": truncated or len(workbook.sheets()) > SPREADSHEET_PREVIEW_MAX_SHEETS,
            })
        workbook.release_resources()
        return sheets
    except Exception as error:  # parser errors are normalized for the API boundary
        raise DomainError("PREVIEW_RENDER_FAILED", "旧版 Excel 文件无法解析为安全预览，请下载原件查看", 503) from error


def spreadsheet_to_preview(data: bytes, media_type: str, filename: str) -> dict[str, object]:
    """Return a bounded, JSON-serializable preview for XLSX, XLS or CSV."""
    suffix = Path(filename).suffix.lower()
    if suffix == ".csv" or media_type == "text/csv":
        sheets = _csv_preview(data)
        file_format = "csv"
    elif suffix == ".xls" or media_type == "application/vnd.ms-excel":
        sheets = _xls_preview(data)
        file_format = "xls"
    elif suffix == ".xlsx" or media_type == (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    ):
        sheets = _xlsx_preview(data)
        file_format = "xlsx"
    else:
        raise DomainError("PREVIEW_RENDER_UNSUPPORTED", "当前文件不支持表格预览", 415)
    return {
        "kind": "spreadsheet",
        "format": file_format,
        "filename": filename,
        "sheets": sheets,
        "limits": {
            "max_rows": SPREADSHEET_PREVIEW_MAX_ROWS,
            "max_columns": SPREADSHEET_PREVIEW_MAX_COLUMNS,
            "max_sheets": SPREADSHEET_PREVIEW_MAX_SHEETS,
        },
    }
